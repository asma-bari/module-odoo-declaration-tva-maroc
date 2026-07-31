import base64
import logging
import xml.etree.ElementTree as ET

from io import BytesIO

from openpyxl import Workbook

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from odoo import api, fields, models
from odoo.exceptions import ValidationError

from odoo.exceptions import ValidationError
_logger = logging.getLogger(__name__)

from openpyxl.styles import Font, Alignment

class VatReturn(models.Model):
    _name = "l10n.ma.vat.return"
    _description = "Déclaration TVA Maroc"
    _inherit = ["mail.thread", "mail.activity.mixin"]
    _rec_name = "name"
    _order = "period_start desc"

    # Informations générales

    name = fields.Char(
        string="Référence",
        required=True,
        default="Nouveau",
        copy=False,
        tracking=True,
    )

    company_id = fields.Many2one(
        "res.company",
        string="Société",
        required=True,
        default=lambda self: self.env.company,
        tracking=True,
    )

    currency_id = fields.Many2one(
        "res.currency",
        related="company_id.currency_id",
        store=True,
        readonly=True,
    )

    period_start = fields.Date(
        string="Date début",
        required=True,
        tracking=True,
    )

    period_end = fields.Date(
        string="Date fin",
        required=True,
        tracking=True,
    )

    regime = fields.Selection(
        [
            ("cash", "Encaissements"),
            ("accrual", "Débits"),
        ],
        string="Régime",
        default="cash",
        required=True,
        tracking=True,
    )

    periodicity = fields.Selection(
        [
            ("monthly", "Mensuelle"),
            ("quarterly", "Trimestrielle"),
        ],
        string="Périodicité",
        default="monthly",
        required=True,
        tracking=True,
    )

    state = fields.Selection(
        [
            ("draft", "Brouillon"),
            ("prepared", "Préparée"),
            ("checked", "Contrôlée"),
            ("validated", "Validée"),
            ("submitted", "Déposée"),
            ("cancel", "Annulée"),
        ],
        string="État",
        default="draft",
        tracking=True,
    )

    # Totaux TVA


    vat_collected = fields.Monetary(
        string="TVA collectée",
        currency_field="currency_id",
        compute="_compute_totals",
        store=True,
    )

    vat_deductible = fields.Monetary(
        string="TVA déductible",
        currency_field="currency_id",
        compute="_compute_totals",
        store=True,
    )

    vat_credit = fields.Monetary(
        string="Crédit TVA",
        currency_field="currency_id",
        compute="_compute_totals",
        store=True,
    )

    vat_due = fields.Monetary(
        string="TVA à payer",
        currency_field="currency_id",
        compute="_compute_totals",
        store=True,
    )


    # Validation

    validation_date = fields.Datetime(
        string="Date validation",
        readonly=True,
        tracking=True,
    )

    validated_by = fields.Many2one(
        "res.users",
        string="Validée par",
        readonly=True,
        tracking=True,
    )


    # Documents

    xml_file = fields.Binary(
        string="Fichier XML",
        attachment=True,
    )

    xml_filename = fields.Char()

    excel_file = fields.Binary(
        string="Fichier Excel",
        attachment=True,
    )

    excel_filename = fields.Char()

    pdf_file = fields.Binary(
        string="Rapport PDF",
        attachment=True,
    )

    pdf_filename = fields.Char()

    note = fields.Text(
        string="Observations",
    )

    active = fields.Boolean(
        default=True,
    )

    # Relations


    line_ids = fields.One2many(
        "l10n.ma.vat.return.line",
        "vat_return_id",
        string="Lignes TVA",
    )

    deduction_line_ids = fields.One2many(
        "l10n.ma.vat.deduction.line",
        "vat_return_id",
        string="Déductions",
    )

    history_ids = fields.One2many(
        "l10n.ma.vat.history",
        "vat_return_id",
        string="Historique",
    )

    # Contraintes


    _sql_constraints = [
        (
            "unique_period_company",
            "unique(company_id,period_start,period_end)",
            "Une déclaration existe déjà pour cette période.",
        ),
    ]

    # Création automatique de la référence


    @api.model_create_multi
    def create(self, vals_list):

        for vals in vals_list:

            if vals.get("name", "Nouveau") == "Nouveau":
                vals["name"] = (
                    self.env["ir.sequence"].next_by_code(
                        "l10n.ma.vat.return"
                    )
                    or "Nouveau"
                )

        return super().create(vals_list)

    # Calcul des totaux


    @api.depends("line_ids.vat_amount", "line_ids.line_type")
    def _compute_totals(self):

        for record in self:

            collected = 0.0
            deductible = 0.0

            for line in record.line_ids:

                if line.line_type in ("sale", "autoliquidation"):
                    collected += line.vat_amount

                elif line.line_type in (
                    "purchase",
                    "deduction",
                    "credit",
                ):
                    deductible += line.vat_amount

            record.vat_collected = collected
            record.vat_deductible = deductible

            difference = collected - deductible

            if difference >= 0:
                record.vat_due = difference
                record.vat_credit = 0.0
            else:
                record.vat_due = 0.0
                record.vat_credit = abs(difference)

    # Workflow

    def action_prepare(self):
        self.ensure_one()

        if self.period_start > self.period_end:
            raise ValidationError(
                "La date de début doit être antérieure à la date de fin."
            )

        # Supprimer les anciennes lignes
        self.line_ids.unlink()
        self.deduction_line_ids.unlink()

        # 1. Recherche des factures clients


        invoices = self.env["account.move"].search([
            ("move_type", "=", "out_invoice"),
            ("state", "=", "posted"),
            ("invoice_date", ">=", self.period_start),
            ("invoice_date", "<=", self.period_end),
        ])

        if not invoices:
            raise ValidationError(
                "Aucune facture client trouvée pour cette période."
            )

        # 2. Création des lignes TVA


        for invoice in invoices:

            for invoice_line in invoice.invoice_line_ids:

                if not invoice_line.tax_ids:
                    continue

                tax = invoice_line.tax_ids[0]

                self.env["l10n.ma.vat.return.line"].create({
                    "vat_return_id": self.id,
                    "sequence": 10,
                    "name": invoice_line.name,
                    "line_type": "sale",
                    "tax_rate": tax.amount,
                    "base_amount": invoice_line.price_subtotal,
                    "vat_amount": (
                        invoice_line.price_total
                        - invoice_line.price_subtotal
                    ),
                })

        # 3. Recherche des factures fournisseurs


        vendor_bills = self.env["account.move"].search([
            ("move_type", "=", "in_invoice"),
            ("state", "=", "posted"),
            ("invoice_date", ">=", self.period_start),
            ("invoice_date", "<=", self.period_end),
        ])

        # 4. Création des lignes de déduction


        for bill in vendor_bills:

            if not bill.partner_id:
                continue

            tax_amount = 0.0

            for line in bill.invoice_line_ids:

                tax_amount += (
                    line.price_total
                    - line.price_subtotal
                )

            self.env["l10n.ma.vat.deduction.line"].create({
                "vat_return_id": self.id,
                "move_id": bill.id,
                "partner_id": bill.partner_id.id,
                "invoice_number": bill.name,
                "invoice_date": bill.invoice_date,
                "amount_untaxed": bill.amount_untaxed,
                "tax_amount": tax_amount,
                "total_amount": bill.amount_total,
            })

        # 5. Changer l'état


        self.write({
            "state": "prepared",
        })

        self._create_history(
            "prepare",
            "La déclaration a été préparée."
        )

    def action_check(self):
        if self.state != "prepared":
            raise ValidationError(
            "Seule une déclaration préparée peut être contrôlée."
        )
        self.write({"state": "checked"})

        self._create_history(
            "check",
            "Déclaration contrôlée.",
        )

        return True

    def action_validate(self):
        if self.state != "checked":
            raise ValidationError(
            "Seule une déclaration contrôlée peut être validée."
        )
        self.write({
            "state": "validated",
            "validation_date": fields.Datetime.now(),
            "validated_by": self.env.user.id,
        })

        self._create_history(
            "validate",
            "Déclaration validée.",
        )

        return True

    def action_submit(self):
        if self.state != "validated":
            raise ValidationError(
            "Seule une déclaration validée peut être déposée."
        )
        self.write({"state": "submitted"})

        self._create_history(
            "submit",
            "Déclaration déposée.",
        )

        return True

    def action_cancel(self):
        self.write({"state": "cancel"})

        self._create_history(
            "cancel",
            "Déclaration annulée.",
        )

        return True

    def action_reset_to_draft(self):
        if self.state == "submitted":
            raise ValidationError(
            "Une déclaration déposée ne peut pas être remise en brouillon."
        )
        self.write({"state": "draft"})

        self._create_history(
            "create",
            "Retour au brouillon.",
        )

        return True

    # Historique


    def _create_history(self, action, description):

        self.env["l10n.ma.vat.history"].create({
            "vat_return_id": self.id,
            "user_id": self.env.user.id,
            "action": action,
            "state": self.state,
            "description": description,
        })

    # Export XML


    def action_export_xml(self):
        self.ensure_one()

        root = ET.Element("vat_return")

        # ==========================
        # Informations générales
        # ==========================

        ET.SubElement(root, "reference").text = self.name or ""
        ET.SubElement(root, "company").text = self.company_id.name or ""
        ET.SubElement(root, "period_start").text = str(self.period_start or "")
        ET.SubElement(root, "period_end").text = str(self.period_end or "")
        ET.SubElement(root, "regime").text = self.regime or ""
        ET.SubElement(root, "periodicity").text = self.periodicity or ""
        ET.SubElement(root, "state").text = self.state or ""

        # Totaux TVA


        totals = ET.SubElement(root, "totals")

        ET.SubElement(
            totals,
            "vat_collected"
        ).text = str(self.vat_collected)

        ET.SubElement(
            totals,
            "vat_deductible"
        ).text = str(self.vat_deductible)

        ET.SubElement(
            totals,
            "vat_credit"
        ).text = str(self.vat_credit)

        ET.SubElement(
            totals,
            "vat_due"
        ).text = str(self.vat_due)

        # Lignes TVA


        lines = ET.SubElement(root, "lines")

        for line in self.line_ids:

            xml_line = ET.SubElement(lines, "line")

            ET.SubElement(xml_line, "sequence").text = str(line.sequence)
            ET.SubElement(xml_line, "name").text = line.name or ""
            ET.SubElement(xml_line, "type").text = line.line_type or ""
            ET.SubElement(xml_line, "tax_rate").text = str(line.tax_rate)
            ET.SubElement(xml_line, "base_amount").text = str(line.base_amount)
            ET.SubElement(xml_line, "vat_amount").text = str(line.vat_amount)
            ET.SubElement(xml_line, "note").text = line.note or ""

        # Déductions TVA


        deductions = ET.SubElement(root, "deductions")

        for deduction in self.deduction_line_ids:

            xml_deduction = ET.SubElement(
                deductions,
                "deduction"
            )

            ET.SubElement(
                xml_deduction,
                "invoice_number"
            ).text = deduction.invoice_number or ""

            ET.SubElement(
                xml_deduction,
                "supplier"
            ).text = deduction.partner_id.name or ""

            ET.SubElement(
                xml_deduction,
                "invoice_date"
            ).text = str(deduction.invoice_date or "")

            ET.SubElement(
                xml_deduction,
                "payment_date"
            ).text = str(deduction.payment_date or "")

            ET.SubElement(
                xml_deduction,
                "amount_untaxed"
            ).text = str(deduction.amount_untaxed)

            ET.SubElement(
                xml_deduction,
                "tax_amount"
            ).text = str(deduction.tax_amount)

            ET.SubElement(
                xml_deduction,
                "total_amount"
            ).text = str(deduction.total_amount)

            ET.SubElement(
                xml_deduction,
                "supplier_ice"
            ).text = deduction.supplier_ice or ""

            ET.SubElement(
                xml_deduction,
                "supplier_if"
            ).text = deduction.supplier_if or ""

            ET.SubElement(
                xml_deduction,
                "supplier_rc"
            ).text = deduction.supplier_rc or ""

            ET.SubElement(
                xml_deduction,
                "payment_method"
            ).text = deduction.payment_method or ""


        # Création du fichier


        xml_content = ET.tostring(
            root,
            encoding="utf-8",
            xml_declaration=True,
        )

        self.write({
            "xml_file": base64.b64encode(xml_content),
            "xml_filename": "%s.xml" % self.name,
        })

        self._create_history(
            "export_xml",
            "Fichier XML généré."
        )

        return True        

    # Export Excel


    def action_export_excel(self):
        self.ensure_one()

        output = BytesIO()

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Déclaration TVA"

        bold_font = Font(bold=True)
        center = Alignment(horizontal="center")

        # Largeur des colonnes

        sheet.column_dimensions["A"].width = 25
        sheet.column_dimensions["B"].width = 25
        sheet.column_dimensions["C"].width = 20
        sheet.column_dimensions["D"].width = 20
        sheet.column_dimensions["E"].width = 20
        sheet.column_dimensions["F"].width = 20

        # Informations générales


        sheet["A1"] = "Déclaration TVA Maroc"
        sheet["A1"].font = Font(size=16, bold=True)
        sheet["A1"].alignment = center

        sheet["A3"] = "Référence"
        sheet["A3"].font = bold_font
        sheet["B3"] = self.name

        sheet["A4"] = "Société"
        sheet["A4"].font = bold_font
        sheet["B4"] = self.company_id.name

        sheet["A5"] = "Date début"
        sheet["A5"].font = bold_font
        sheet["B5"] = str(self.period_start)

        sheet["A6"] = "Date fin"
        sheet["A6"].font = bold_font
        sheet["B6"] = str(self.period_end)

        sheet["A7"] = "Régime"
        sheet["A7"].font = bold_font
        sheet["B7"] = self.regime

        sheet["A8"] = "Périodicité"
        sheet["A8"].font = bold_font
        sheet["B8"] = self.periodicity

        sheet["A9"] = "État"
        sheet["A9"].font = bold_font
        sheet["B9"] = self.state

        # Totaux TVA


        sheet["A11"] = "TVA collectée"
        sheet["A11"].font = bold_font
        sheet["B11"] = self.vat_collected

        sheet["A12"] = "TVA déductible"
        sheet["A12"].font = bold_font
        sheet["B12"] = self.vat_deductible

        sheet["A13"] = "Crédit TVA"
        sheet["A13"].font = bold_font
        sheet["B13"] = self.vat_credit

        sheet["A14"] = "TVA à payer"
        sheet["A14"].font = bold_font
        sheet["B14"] = self.vat_due

        # Lignes TVA


        row = 17

        sheet.cell(row=row, column=1).value = "Rubrique"
        sheet.cell(row=row, column=2).value = "Type"
        sheet.cell(row=row, column=3).value = "Taux TVA"
        sheet.cell(row=row, column=4).value = "Montant HT"
        sheet.cell(row=row, column=5).value = "Montant TVA"

        for col in range(1, 6):
            sheet.cell(row=row, column=col).font = bold_font
            sheet.cell(row=row, column=col).alignment = center

        row += 1

        for line in self.line_ids:

            sheet.cell(row=row, column=1).value = line.name
            sheet.cell(row=row, column=2).value = line.line_type
            sheet.cell(row=row, column=3).value = line.tax_rate
            sheet.cell(row=row, column=4).value = line.base_amount
            sheet.cell(row=row, column=5).value = line.vat_amount

            row += 1

        # Déductions


        row += 2

        sheet.cell(row=row, column=1).value = "Déductions"
        sheet.cell(row=row, column=1).font = Font(size=14, bold=True)

        row += 1

        sheet.cell(row=row, column=1).value = "Facture"
        sheet.cell(row=row, column=2).value = "Fournisseur"
        sheet.cell(row=row, column=3).value = "Date"
        sheet.cell(row=row, column=4).value = "Montant HT"
        sheet.cell(row=row, column=5).value = "Montant TVA"
        sheet.cell(row=row, column=6).value = "Montant TTC"

        for col in range(1, 7):
            sheet.cell(row=row, column=col).font = bold_font
            sheet.cell(row=row, column=col).alignment = center

        row += 1

        for deduction in self.deduction_line_ids:

            sheet.cell(row=row, column=1).value = deduction.invoice_number
            sheet.cell(row=row, column=2).value = deduction.partner_id.name
            sheet.cell(row=row, column=3).value = str(deduction.invoice_date)
            sheet.cell(row=row, column=4).value = deduction.amount_untaxed
            sheet.cell(row=row, column=5).value = deduction.tax_amount
            sheet.cell(row=row, column=6).value = deduction.total_amount

            row += 1

        workbook.save(output)
        output.seek(0)

        self.write({
            "excel_file": base64.b64encode(output.read()),
            "excel_filename": f"{self.name}.xlsx",
        })

        self._create_history(
            "export_excel",
            "Fichier Excel généré."
        )

        return True
      

    # Export PDF


    def action_export_pdf(self):
        self.ensure_one()

        buffer = BytesIO()

        pdf = canvas.Canvas(buffer, pagesize=A4)

        width, height = A4

        y = height - 50


        # Titre
 

        pdf.setTitle("Déclaration TVA Maroc")


        # Titre principal


        pdf.setFont("Helvetica-Bold", 20)
        pdf.drawCentredString(
            width / 2,
            y,
            "DÉCLARATION TVA MAROC"
        )

        y -= 15

        pdf.setLineWidth(1)
        pdf.line(40, y, width - 40, y)

        y -= 35

        # Informations générales
 

        pdf.setFont("Helvetica-Bold", 12)
        pdf.drawString(50, y, "Référence :")
        pdf.setFont("Helvetica", 12)
        pdf.drawString(170, y, self.name)

        y -= 22

        pdf.setFont("Helvetica-Bold", 12)
        pdf.drawString(50, y, "Société :")
        pdf.setFont("Helvetica", 12)
        pdf.drawString(170, y, self.company_id.name)

        y -= 22

        pdf.setFont("Helvetica-Bold", 12)
        pdf.drawString(50, y, "Période :")
        pdf.setFont("Helvetica", 12)
        pdf.drawString(
            170,
            y,
            f"{self.period_start} → {self.period_end}"
        )

        y -= 22

        pdf.setFont("Helvetica-Bold", 12)
        pdf.drawString(50, y, "Régime :")
        pdf.setFont("Helvetica", 12)
        pdf.drawString(170, y, self.regime)

        y -= 22

        pdf.setFont("Helvetica-Bold", 12)
        pdf.drawString(50, y, "Périodicité :")
        pdf.setFont("Helvetica", 12)
        pdf.drawString(170, y, self.periodicity)

        y -= 22

        pdf.setFont("Helvetica-Bold", 12)
        pdf.drawString(50, y, "État :")
        pdf.setFont("Helvetica", 12)
        pdf.drawString(170, y, self.state)

        y -= 30

        pdf.setFont("Helvetica-Bold", 12)
        pdf.drawString(50, y, "Date de génération :")

        pdf.setFont("Helvetica", 12)
        pdf.drawString(
            170,
            y,
            fields.Datetime.now().strftime("%d/%m/%Y %H:%M")
        )

        y -= 30       

        pdf.line(40, y, width - 40, y)

        y -= 30     

        pdf.setFont("Helvetica-Bold", 12)
        pdf.drawString(50, y, "Généré par :")

        pdf.setFont("Helvetica", 12)
        pdf.drawString(
            170,
            y,
            self.env.user.name
        )

        y -= 30   

        # Totaux TVA



        pdf.setFont("Helvetica-Bold", 14)
        pdf.drawString(50, y, "TOTAUX TVA")

        y -= 20

        pdf.line(40, y, width - 40, y)

        y -= 25

        pdf.setFont("Helvetica-Bold", 11)

        pdf.drawString(60, y, "TVA collectée")
        pdf.drawRightString(width - 60, y, f"{self.vat_collected:.2f} DH")

        y -= 22

        pdf.drawString(60, y, "TVA déductible")
        pdf.drawRightString(width - 60, y, f"{self.vat_deductible:.2f} DH")

        y -= 22

        pdf.drawString(60, y, "Crédit TVA")
        pdf.drawRightString(width - 60, y, f"{self.vat_credit:.2f} DH")

        y -= 22

        pdf.drawString(60, y, "TVA à payer")
        pdf.drawRightString(width - 60, y, f"{self.vat_due:.2f} DH")

        y -= 15

        pdf.line(40, y, width - 40, y)

        y -= 35

        # Lignes TVA

        pdf.setFont("Helvetica-Bold", 14)
        pdf.drawString(50, y, "LIGNES TVA")

        y -= 25

        # En-tête du tableau
        pdf.setFont("Helvetica-Bold", 10)

        pdf.drawString(50, y, "Rubrique")
        pdf.drawString(180, y, "Type")
        pdf.drawString(270, y, "Taux")
        pdf.drawString(340, y, "HT")
        pdf.drawString(460, y, "TVA")

        y -= 10
        pdf.line(40, y, width - 40, y)
        y -= 20

        pdf.setFont("Helvetica", 10)

        for line in self.line_ids:

            pdf.drawString(50, y, line.name or "")
            pdf.drawString(180, y, line.line_type or "")
            pdf.drawString(270, y, f"{line.tax_rate:.2f}%")
            pdf.drawRightString(430, y, f"{line.base_amount:.2f}")
            pdf.drawRightString(550, y, f"{line.vat_amount:.2f}")

            y -= 18

            if y < 80:
                pdf.showPage()

                y = height - 50

                pdf.setFont("Helvetica-Bold", 14)
                pdf.drawString(50, y, "LIGNES TVA")

                y -= 25

                pdf.setFont("Helvetica-Bold", 10)

                pdf.drawString(50, y, "Rubrique")
                pdf.drawString(180, y, "Type")
                pdf.drawString(270, y, "Taux")
                pdf.drawString(340, y, "HT")
                pdf.drawString(460, y, "TVA")

                y -= 10
                pdf.line(40, y, width - 40, y)
                y -= 20

                pdf.setFont("Helvetica", 10)

        y -= 20

        # Déductions
  
        pdf.setFont("Helvetica-Bold", 14)
        pdf.drawString(50, y, "DÉDUCTIONS")

        y -= 25

        # En-tête du tableau

        pdf.setFont("Helvetica-Bold", 10)

        pdf.drawString(50, y, "Facture")
        pdf.drawString(130, y, "Fournisseur")
        pdf.drawString(280, y, "Date")
        pdf.drawString(360, y, "HT")
        pdf.drawString(450, y, "TVA")
        pdf.drawString(530, y, "TTC")

        y -= 10

        pdf.line(40, y, width - 40, y)

        y -= 20

        pdf.setFont("Helvetica", 9)

        for deduction in self.deduction_line_ids:

            pdf.drawString(
                50,
                y,
                deduction.invoice_number or ""
            )

            pdf.drawString(
                130,
                y,
                deduction.partner_id.name or ""
            )

            pdf.drawString(
                280,
                y,
                str(deduction.invoice_date or "")
            )

            pdf.drawRightString(
                430,
                y,
                f"{deduction.amount_untaxed:.2f}"
            )

            pdf.drawRightString(
                500,
                y,
                f"{deduction.tax_amount:.2f}"
            )

            pdf.drawRightString(
                570,
                y,
                f"{deduction.total_amount:.2f}"
            )

            y -= 18

            if y < 80:

                pdf.showPage()

                y = height - 50

                pdf.setFont("Helvetica-Bold", 14)
                pdf.drawString(50, y, "DÉDUCTIONS")

                y -= 25

                pdf.setFont("Helvetica-Bold", 10)

                pdf.drawString(50, y, "Facture")
                pdf.drawString(130, y, "Fournisseur")
                pdf.drawString(280, y, "Date")
                pdf.drawString(360, y, "HT")
                pdf.drawString(450, y, "TVA")
                pdf.drawString(530, y, "TTC")

                y -= 10

                pdf.line(40, y, width - 40, y)

                y -= 20

                pdf.setFont("Helvetica", 9)

        y -= 30

        pdf.setFont("Helvetica-Oblique", 9)

        pdf.drawCentredString(
            width / 2,
            20,
            "Déclaration TVA Maroc - Générée automatiquement par Odoo"
        )

        pdf.save()

        pdf_data = buffer.getvalue()

        buffer.close()

        self.write({
            "pdf_file": base64.b64encode(pdf_data),
            "pdf_filename": f"{self.name}.pdf",
        })

        self._create_history(
            "export_pdf",
            "Rapport PDF généré."
        )

        return True    
    @api.constrains("company_id", "period_start", "period_end")
    def _check_duplicate_period(self):

        for record in self:

            duplicate = self.search([
                ("id", "!=", record.id),
                ("company_id", "=", record.company_id.id),
                ("period_start", "=", record.period_start),
                ("period_end", "=", record.period_end),
            ], limit=1)

            if duplicate:
                raise ValidationError(
                    "Une déclaration TVA existe déjà pour cette société et cette période."
                )    